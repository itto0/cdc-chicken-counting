from datetime import datetime, timezone
from io import BytesIO
import openpyxl
import pytest
from reportlab.platypus import Paragraph, Table

from services.api.src.services.exporter import ExcelReportGenerator
from services.api.src.services.pdf_exporter import PDFReportGenerator
import services.api.src.services.exporter as exporter_module
import services.api.src.services.pdf_exporter as pdf_exporter_module
import shared.time_utils as shared_time_utils


def test_shared_time_utils_imported():
    """Verify that exporter and pdf_exporter import standardized time utilities from shared.time_utils."""
    assert hasattr(exporter_module, "format_jakarta_display")
    assert exporter_module.format_jakarta_display is shared_time_utils.format_jakarta_display

    assert hasattr(pdf_exporter_module, "format_jakarta_display")
    assert pdf_exporter_module.format_jakarta_display is shared_time_utils.format_jakarta_display
    assert hasattr(pdf_exporter_module, "now_jakarta")
    assert pdf_exporter_module.now_jakarta is shared_time_utils.now_jakarta


def test_excel_exporter_utc_to_wib_conversion():
    """Excel generator converts UTC datetimes (aware & naive) and ISO strings to Jakarta WIB."""
    generator = ExcelReportGenerator()

    # 1. Aware UTC datetime: 03:00 UTC -> 10:00 WIB
    records_aware = [
        {
            "window_start": datetime(2026, 4, 14, 3, 0, 0, tzinfo=timezone.utc),
            "count": 100,
            "rate": 50.0,
            "status": "NORMAL",
        }
    ]
    xlsx_bytes = generator.generate_line_report(line_id="line-aware", records=records_aware)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][0] == "2026-04-14 10:00:00"

    # 2. Naive UTC datetime: 03:00 -> 10:00 WIB
    records_naive = [
        {
            "window_start": datetime(2026, 4, 14, 3, 0, 0),
            "count": 100,
            "rate": 50.0,
            "status": "NORMAL",
        }
    ]
    xlsx_bytes_naive = generator.generate_line_report(line_id="line-naive", records=records_naive)
    wb_naive = openpyxl.load_workbook(BytesIO(xlsx_bytes_naive))
    ws_naive = wb_naive.active
    rows_naive = list(ws_naive.iter_rows(values_only=True))
    assert rows_naive[1][0] == "2026-04-14 10:00:00"

    # 3. ISO formatted string with Z: 03:00Z -> 10:00 WIB
    records_iso = [
        {
            "window_start": "2026-04-14T03:00:00Z",
            "count": 100,
            "rate": 50.0,
            "status": "NORMAL",
        }
    ]
    xlsx_bytes_iso = generator.generate_line_report(line_id="line-iso", records=records_iso)
    wb_iso = openpyxl.load_workbook(BytesIO(xlsx_bytes_iso))
    ws_iso = wb_iso.active
    rows_iso = list(ws_iso.iter_rows(values_only=True))
    assert rows_iso[1][0] == "2026-04-14 10:00:00"


def test_pdf_exporter_elements_and_wib_timezone():
    """PDF generator includes WIB in headers, table columns, and converts UTC timestamps."""
    generator = PDFReportGenerator()

    # Aware UTC: 03:00 UTC -> 10:00 WIB
    utc_dt = datetime(2026, 4, 14, 3, 0, 0, tzinfo=timezone.utc)
    records = [
        {
            "window_start": utc_dt,
            "line_id": "line-1",
            "count": 120,
            "rate_per_min": 60.0,
            "status": "NORMAL",
        }
    ]

    elements = generator._build_elements(line_id="line-1", records=records)

    # 1. Header paragraph includes "WIB"
    header_paragraphs = [e for e in elements if isinstance(e, Paragraph) and "Generated at:" in getattr(e, "text", "")]
    assert len(header_paragraphs) == 1
    header_text = header_paragraphs[0].text
    assert "WIB" in header_text

    # 2. Table column header has "Interval Start (WIB)"
    tables = [e for e in elements if isinstance(e, Table)]
    assert len(tables) >= 2
    records_table = tables[1]
    headers = records_table._cellvalues[0]
    assert headers[0] == "Interval Start (WIB)"

    # 3. Formatted interval timestamp matches Jakarta local time: "2026-04-14 10:00"
    row1 = records_table._cellvalues[1]
    assert row1[0] == "2026-04-14 10:00"

    # 4. Naive UTC datetime also converted to Jakarta local time
    records_naive = [
        {
            "window_start": datetime(2026, 4, 14, 3, 0, 0),
            "line_id": "line-1",
            "count": 120,
            "rate_per_min": 60.0,
            "status": "NORMAL",
        }
    ]
    elements_naive = generator._build_elements(line_id="line-1", records=records_naive)
    tables_naive = [e for e in elements_naive if isinstance(e, Table)]
    assert tables_naive[1]._cellvalues[1][0] == "2026-04-14 10:00"

    # 5. ISO string with Z also converted to Jakarta local time
    records_iso = [
        {
            "window_start": "2026-04-14T03:00:00Z",
            "line_id": "line-1",
            "count": 120,
            "rate_per_min": 60.0,
            "status": "NORMAL",
        }
    ]
    elements_iso = generator._build_elements(line_id="line-1", records=records_iso)
    tables_iso = [e for e in elements_iso if isinstance(e, Table)]
    assert tables_iso[1]._cellvalues[1][0] == "2026-04-14 10:00"


def test_pdf_exporter_generates_valid_pdf_bytes():
    """generate_line_report returns valid PDF bytes starting with %PDF."""
    generator = PDFReportGenerator()
    records = [
        {
            "window_start": datetime(2026, 4, 14, 3, 0, 0, tzinfo=timezone.utc),
            "count": 42,
            "rate_per_min": 42.0,
            "status": "NORMAL",
        }
    ]
    pdf_bytes = generator.generate_line_report("line-1", records)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
