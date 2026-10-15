from datetime import datetime, timedelta, timezone
from io import BytesIO
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from services.api.src.db.models import Base, CountAggregate
from services.api.src.db.session import get_db
from services.api.src.main import app
from services.api.src.services.exporter import ExcelReportGenerator


@pytest.fixture
def setup_test_db(tmp_path):
    test_db_file = tmp_path / "test_api_exporter.db"
    test_engine = create_engine(
        f"sqlite:///{test_db_file}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=test_engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    yield TestingSessionLocal
    app.dependency_overrides.clear()
    test_engine.dispose()


def test_excel_report_generation_unit():
    records = [
        {"ts_wib": "2026-10-15 08:00:00", "count": 120, "rate": 120.0, "status": "NORMAL"},
        {"ts_wib": "2026-10-15 08:01:00", "count": 115, "rate": 115.0, "status": "NORMAL"},
    ]
    generator = ExcelReportGenerator()
    xlsx_bytes = generator.generate_line_report(line_id="lineA", records=records)

    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 1000

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    assert "Report lineA" in wb.sheetnames
    ws = wb["Report lineA"]

    # Verify header
    headers = [cell.value for cell in ws[1]]
    assert headers == ["Timestamp (WIB)", "Line", "Count", "Rate (birds/min)", "Status"]

    # Verify data rows
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[1] == ("2026-10-15 08:00:00", "lineA", 120, 120.0, "NORMAL")
    assert rows[2] == ("2026-10-15 08:01:00", "lineA", 115, 115.0, "NORMAL")


def test_excel_report_generation_empty_records():
    generator = ExcelReportGenerator()
    xlsx_bytes = generator.generate_line_report(line_id="lineEmpty", records=[])

    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 1000

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Report lineEmpty"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 1
    assert rows[0] == ("Timestamp (WIB)", "Line", "Count", "Rate (birds/min)", "Status")


def test_excel_report_generator_wib_datetime_conversion():
    generator = ExcelReportGenerator()
    # UTC datetime 01:00:00 should become 08:00:00 WIB (UTC+7)
    utc_time = datetime(2026, 10, 15, 1, 0, 0, tzinfo=timezone.utc)
    records = [
        {"window_start": utc_time, "count": 50, "rate_per_min": 50.0, "status": "WARNING"}
    ]
    xlsx_bytes = generator.generate_line_report(line_id="lineA", records=records)

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Report lineA"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][0] == "2026-10-15 08:00:00"
    assert rows[1][1] == "lineA"
    assert rows[1][2] == 50
    assert rows[1][3] == 50.0
    assert rows[1][4] == "WARNING"


def test_export_endpoint_default_empty(setup_test_db):
    client = TestClient(app)
    resp = client.get("/api/v1/reports/export?line_id=lineA")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "report_lineA.xlsx" in resp.headers["content-disposition"]
    assert "attachment" in resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb["Report lineA"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 1
    assert rows[0] == ("Timestamp (WIB)", "Line", "Count", "Rate (birds/min)", "Status")


def test_export_endpoint_with_records(setup_test_db):
    with setup_test_db() as db:
        t1 = datetime(2026, 10, 15, 1, 0, 0, tzinfo=timezone.utc)
        t2 = datetime(2026, 10, 15, 1, 1, 0, tzinfo=timezone.utc)
        t_other = datetime(2026, 10, 15, 1, 0, 0, tzinfo=timezone.utc)

        db.add(CountAggregate(
            line_id="lineA",
            window_start=t1,
            window_end=t1 + timedelta(minutes=1),
            count=100,
            rate_per_min=100.0,
            confidence_avg=0.95,
            status="NORMAL"
        ))
        db.add(CountAggregate(
            line_id="lineA",
            window_start=t2,
            window_end=t2 + timedelta(minutes=1),
            count=95,
            rate_per_min=95.0,
            confidence_avg=0.94,
            status="WARNING"
        ))
        db.add(CountAggregate(
            line_id="lineB",
            window_start=t_other,
            window_end=t_other + timedelta(minutes=1),
            count=200,
            rate_per_min=200.0,
            confidence_avg=0.98,
            status="NORMAL"
        ))
        db.commit()

    client = TestClient(app)
    resp = client.get("/api/v1/reports/export?line_id=lineA")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb["Report lineA"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[1] == ("2026-10-15 08:00:00", "lineA", 100, 100.0, "NORMAL")
    assert rows[2] == ("2026-10-15 08:01:00", "lineA", 95, 95.0, "WARNING")


def test_export_endpoint_with_time_range_filter(setup_test_db):
    with setup_test_db() as db:
        t1 = datetime(2026, 10, 15, 1, 0, 0, tzinfo=timezone.utc)
        t2 = datetime(2026, 10, 15, 2, 0, 0, tzinfo=timezone.utc)
        t3 = datetime(2026, 10, 15, 3, 0, 0, tzinfo=timezone.utc)

        for t, cnt in [(t1, 10), (t2, 20), (t3, 30)]:
            db.add(CountAggregate(
                line_id="lineA",
                window_start=t,
                window_end=t + timedelta(minutes=1),
                count=cnt,
                rate_per_min=float(cnt),
                confidence_avg=0.95,
                status="NORMAL"
            ))
        db.commit()

    client = TestClient(app)
    # Filter only around t2 (02:00:00 UTC)
    start = "2026-10-15T01:30:00Z"
    end = "2026-10-15T02:30:00Z"
    resp = client.get(f"/api/v1/reports/export?line_id=lineA&start_time={start}&end_time={end}")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb["Report lineA"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][2] == 20


def test_export_endpoint_unsupported_format():
    client = TestClient(app)
    resp = client.get("/api/v1/reports/export?line_id=lineA&format=csv")
    assert resp.status_code == 400
    assert "Unsupported format" in resp.json()["detail"]


def test_export_endpoint_invalid_datetime_format():
    client = TestClient(app)
    resp = client.get("/api/v1/reports/export?line_id=lineA&start_time=invalid-date")
    assert resp.status_code == 400
    assert "Invalid datetime" in resp.json()["detail"]
