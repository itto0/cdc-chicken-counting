from datetime import datetime
from io import BytesIO
import re
from typing import Any
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared.time_utils import format_jakarta_display


class ExcelReportGenerator:
    """Generates professional Excel reports (.xlsx) for poultry production lines."""

    HEADER_COLUMNS = [
        "Timestamp (WIB)",
        "Line",
        "Count",
        "Rate (birds/min)",
        "Status",
    ]

    def _format_timestamp(self, ts_val: Any) -> str:
        if isinstance(ts_val, datetime):
            return format_jakarta_display(ts_val)
        if isinstance(ts_val, str) and ("T" in ts_val or ts_val.endswith("Z")):
            try:
                clean_ts = ts_val.strip().replace("Z", "+00:00")
                dt = datetime.fromisoformat(clean_ts)
                return format_jakarta_display(dt)
            except Exception:
                return str(ts_val)
        if ts_val is not None:
            return str(ts_val)
        return ""

    def generate_line_report(self, line_id: str, records: list[dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        clean_line_id = re.sub(r"[\\/*?:\[\]]", "_", str(line_id))
        safe_title = f"Report {clean_line_id}"[:31]
        ws.title = safe_title

        # Header styling: Corporate dark navy fill, white bold text
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        thin_side = Side(border_style="thin", color="D3D3D3")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.append(self.HEADER_COLUMNS)
        ws.row_dimensions[1].height = 24

        for col_idx in range(1, len(self.HEADER_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        data_alignments = [
            Alignment(horizontal="center", vertical="center"),  # Timestamp
            Alignment(horizontal="center", vertical="center"),  # Line
            Alignment(horizontal="right", vertical="center"),   # Count
            Alignment(horizontal="right", vertical="center"),   # Rate
            Alignment(horizontal="center", vertical="center"),  # Status
        ]

        for row_idx, r in enumerate(records, start=2):
            ts_val = (
                r.get("ts_wib")
                or r.get("window_start")
                or r.get("timestamp")
                or r.get("ts_utc")
            )
            ts_str = self._format_timestamp(ts_val)
            cnt = r.get("count", 0)
            if cnt is None:
                cnt = 0
            rate = r.get("rate") if "rate" in r else r.get("rate_per_min", 0.0)
            if rate is None:
                rate = 0.0
            status = r.get("status") or "NORMAL"

            row_line_id = r.get("line_id") or line_id
            row_values = [ts_str, row_line_id, cnt, rate, status]
            ws.append(row_values)
            ws.row_dimensions[row_idx].height = 20

            for col_idx, align in enumerate(data_alignments, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = align
                cell.border = thin_border
                if col_idx == 4 and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0"

        # Auto-adjust column widths
        for col_idx in range(1, len(self.HEADER_COLUMNS) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(ws.cell(row=r, column=col_idx).value or ""))
                for r in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
